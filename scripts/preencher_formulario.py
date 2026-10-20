"""Preenche o formulário da atividade (PIT_atividade_8.docx) sem alterar sua estrutura.

Uso: python scripts/preencher_formulario.py [--repo URL] [--app URL] [--video URL]
Links não informados ficam como [PENDENTE]. Os 5 testes dos colegas ficam [PENDENTE]
até serem preenchidos com resultados reais.
"""
import argparse
import copy
from pathlib import Path

import docx

RAIZ = Path(__file__).resolve().parents[1]
MODELO = RAIZ / "entrega" / "modelo" / "PIT_atividade_8.docx"
SAIDA = RAIZ / "entrega" / "PIT_atividade_8_preenchido.docx"
PENDENTE_TESTE = "[PENDENTE: aguardando teste real]"


def run_como(paragrafo, texto, modelo_run=None, negrito=False):
    """Acrescenta um run copiando a formatação (fonte Arial) de um run de referência."""
    r = paragrafo.add_run(texto)
    if modelo_run is not None and modelo_run._r.rPr is not None:
        r._r.insert(0, copy.deepcopy(modelo_run._r.rPr))
    r.bold = negrito
    r.italic = False
    return r


def escrever_celula(linha, texto, col=1):
    """Escreve na célula de resposta usando a fonte do rótulo da mesma linha, sem o negrito."""
    celula = linha.cells[col]
    ref = next((r for p_ in linha.cells[0].paragraphs for r in p_.runs if r.text.strip()), None)
    p = celula.paragraphs[0]
    for r in p.runs:
        r.text = ""
    run_como(p, texto, ref, negrito=False)


def acrescentar_na_celula(celula, texto):
    """Mantém o rótulo da célula e escreve a resposta logo depois dele."""
    p = celula.paragraphs[-1]
    ref = p.runs[0] if p.runs else None
    run_como(p, " " + texto, ref, negrito=False)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", default="[PENDENTE]")
    ap.add_argument("--app", default="[PENDENTE]")
    ap.add_argument("--video", default="[PENDENTE]")
    a = ap.parse_args()

    d = docx.Document(MODELO)
    ps = d.paragraphs

    # Identificação (regiões editáveis "Digite aqui")
    for p, valor in ((ps[3], "Davy Gabriel Stresser Pereira"), (ps[4], "36963089")):
        alvo = next(r for r in p.runs if r.text == "Digite aqui")
        alvo.text = valor

    # Documentação: link do repositório
    p_link = next(p for p in ps if p.text.startswith("Link do repositório"))
    run_como(p_link, a.repo, p_link.runs[-1])
    p_obs = ps[ps.index(p_link) + 1]
    run_como(p_obs, "Documentação em docs/situacao-1 (escopo, UML, IHC, modelo de dados e dicionário), "
                    "docs/situacao-2 e docs/situacao-3.", p_link.runs[-1])

    # Codificação
    t = d.tables[0]
    respostas = ["Python (Flask)", "PostgreSQL", "Render (com domínio via Cloudflare)", "Web"]
    for i, texto in enumerate(respostas):
        escrever_celula(t.rows[i], texto)
    modo = t.rows[4].cells[1]
    for p in modo.paragraphs:
        for r in p.runs:
            if "Tradicional" in p.text and "( )" in r.text:
                r.text = r.text.replace("( )", "(X)")
    assert "(X) Tradicional" in modo.text and modo.text.count("(X)") == 1, modo.text
    escrever_celula(t.rows[5], a.repo)
    escrever_celula(t.rows[6], a.app)
    escrever_celula(t.rows[7], a.video)

    # Testes da solução: 5 tabelas de colegas
    for tabela in d.tables[1:6]:
        acrescentar_na_celula(tabela.rows[0].cells[0], PENDENTE_TESTE)
        acrescentar_na_celula(tabela.rows[0].cells[1], PENDENTE_TESTE)
        for linha in tabela.rows[1:4]:
            acrescentar_na_celula(linha.cells[0], PENDENTE_TESTE)

    # Laudo de qualidade: referências (usa os parágrafos vazios já existentes)
    i_laudo = next(i for i, p in enumerate(ps) if p.text.startswith("Insira a seguir o laudo"))
    ref = ps[i_laudo].runs[0]
    linhas = [
        "Laudo completo: docs/situacao-3/laudo-qualidade.md (no repositório indicado acima).",
        "Evidências (prints das telas): docs/evidencias/. Plano de testes de verificação e validação: "
        "docs/situacao-3/plano-testes.md. Relatório de testes automatizados: docs/situacao-2/relatorio-testes.md.",
        "Resumo da versão 1.0.0: 87 testes automatizados aprovados, cobertura de 97%; 4 defeitos encontrados na "
        "verificação interna (D-01 a D-04), todos corrigidos. Defeitos dos testes dos colegas: " + PENDENTE_TESTE + ".",
        "Correções após os testes: docs/situacao-3/changelog-melhorias.md.",
    ]
    vazios = [p for p in ps[i_laudo + 1:] if not p.text.strip()][:len(linhas)]
    proximo_titulo = next(i for i, p in enumerate(ps) if p.text.startswith("Vídeo da Solução"))
    assert all(ps.index(p) < proximo_titulo for p in vazios), "parágrafos vazios insuficientes antes do vídeo"
    for p, texto in zip(vazios, linhas):
        run_como(p, texto, ref)

    # Vídeo
    escrever_celula(d.tables[6].rows[0], f"{a.video} (roteiro: docs/situacao-3/roteiro-video.md)")

    SAIDA.parent.mkdir(exist_ok=True)
    d.save(SAIDA)
    print("Formulário salvo em", SAIDA.relative_to(RAIZ))


if __name__ == "__main__":
    main()
